# -*- coding: utf-8 -*-
"""Exportación del calendario a Excel, CSV y JSON."""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path


def columnas(registros: list[dict[str, str]]) -> list[str]:
    """Unión de todas las columnas, en el orden en que aparecen."""
    orden: list[str] = []
    for registro in registros:
        for clave in registro:
            if clave not in orden:
                orden.append(clave)
    # Las columnas con fecha primero, para que la planilla sea legible.
    posiciones = {clave: indice for indice, clave in enumerate(orden)}
    orden.sort(key=lambda k: (0 if "FECHA" in k.upper() else 1, posiciones[k]))
    return orden


def _asegurar_carpeta(ruta: Path) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    return ruta


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def a_excel(registros: list[dict[str, str]], ruta: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Falta openpyxl. Instálalo con: pip install openpyxl"
        ) from exc

    _asegurar_carpeta(ruta)
    encabezados = columnas(registros)

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Calendario"

    relleno = PatternFill("solid", fgColor="1F4E79")
    negrita = Font(bold=True, color="FFFFFF")

    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.fill = relleno
        celda.font = negrita
        celda.alignment = Alignment(horizontal="center", vertical="center")

    for registro in registros:
        hoja.append([registro.get(c, "") for c in encabezados])

    for posicion, nombre in enumerate(encabezados, start=1):
        ancho = max(
            [len(str(nombre))] + [len(str(r.get(nombre, ""))) for r in registros] or [10]
        )
        hoja.column_dimensions[get_column_letter(posicion)].width = min(max(ancho + 2, 12), 60)

    hoja.freeze_panes = "A2"
    if registros:
        hoja.auto_filter.ref = (
            f"A1:{get_column_letter(len(encabezados))}{len(registros) + 1}"
        )
    libro.save(ruta)
    return ruta


# --------------------------------------------------------------------------
# CSV (separador «;», que es el que usa Excel en Chile)
# --------------------------------------------------------------------------
def a_csv(registros: list[dict[str, str]], ruta: Path) -> Path:
    _asegurar_carpeta(ruta)
    encabezados = columnas(registros)
    with ruta.open("w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.DictWriter(
            archivo, fieldnames=encabezados, delimiter=";", extrasaction="ignore"
        )
        escritor.writeheader()
        for registro in registros:
            escritor.writerow({c: registro.get(c, "") for c in encabezados})
    return ruta


# --------------------------------------------------------------------------
# JSON
# --------------------------------------------------------------------------
def a_json(registros: list[dict[str, str]], ruta: Path, meta: dict | None = None) -> Path:
    _asegurar_carpeta(ruta)
    documento = {
        "generado_en": datetime.now().isoformat(timespec="seconds"),
        "total": len(registros),
        **(meta or {}),
        "registros": registros,
    }
    ruta.write_text(json.dumps(documento, ensure_ascii=False, indent=2), encoding="utf-8")
    return ruta
